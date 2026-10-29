"""Persistent product and Docker sheets, with one column per collection run."""
from collections import Counter
import os
import tempfile

from openpyxl import load_workbook
import xlsxwriter
from xlsxwriter.utility import xl_rowcol_to_cell

PRODUCTS = {'Desktop Editors': 'ONLYOFFICE/DesktopEditors',
            'Document Server': 'ONLYOFFICE/DocumentServer'}
GROUPS = ('Installations', 'Updates', 'Add-ons')
PLATFORMS = ('Windows', 'Linux', 'macOS')
SCHEMA = 'onlyoffice-download-history-v1'
METADATA = '_history'
DOCKER_REPOSITORY = 'onlyoffice/documentserver'


class HistoryError(ValueError):
    pass


def empty_history():
    return {'runs': [], 'releases': {}, 'assets': {}, 'docker': []}


def run_label(existing, collected_at):
    base = collected_at.strftime('%Y-%m-%d')
    if base not in existing:
        return base
    base = collected_at.strftime('%Y-%m-%d %H:%M:%S')
    result, index = base, 2
    while result in existing:
        result, index = f'{base} ({index})', index + 1
    return result


def record_release(state, product, release):
    key = (product, str(release['id']))
    state['releases'][key] = {'tag': release['tag_name'], 'published': release['published_at']}
    return key


def add_run(state, products, label, docker_count=None):
    if label in state['runs']:
        raise ValueError('Duplicate collection timestamp')
    if docker_count is not None and (type(docker_count) is not int or docker_count < 0):
        raise ValueError('Invalid Docker pull count')
    state['docker'].append(docker_count)
    for entry in state['assets'].values():
        entry['values'].append(None)
    state['runs'].append(label)
    for product, releases in products.items():
        for release, assets in releases:
            release_key = record_release(state, product, release)
            for asset in assets:
                key = (*release_key, str(asset['id']))
                entry = state['assets'].setdefault(key, {'values': [None] * len(state['runs'])})
                entry.update(name=asset['name'], group=asset['category'][0], platform=asset['category'][1])
                entry['values'][-1] = asset['download_count']


def read_metadata(sheet):
    if sheet['A1'].value != SCHEMA:
        raise ValueError('Unsupported history schema')
    state = empty_history()
    state['runs'] = [cell.value for cell in sheet[2][9:]]
    if not state['runs'] or not all(isinstance(value, str) and value for value in state['runs']):
        raise ValueError('Invalid history timestamps')
    if len(set(state['runs'])) != len(state['runs']):
        raise ValueError('Duplicate history timestamps')
    state['docker'] = [None] * len(state['runs'])
    docker_seen = False
    for row in sheet.iter_rows(min_row=3, values_only=True):
        kind, product, release_id, asset_id, tag, published, name, group, platform = row[:9]
        if kind == 'docker':
            values = list(row[9:9 + len(state['runs'])])
            if docker_seen or product != 'Docker Docs' or name != DOCKER_REPOSITORY:
                raise ValueError('Invalid or duplicate Docker history')
            if not all(value is None or type(value) is int and value >= 0 for value in values):
                raise ValueError('Invalid historical Docker pull count')
            state['docker'] = values
            docker_seen = True
            continue
        if product not in PRODUCTS:
            raise ValueError('Unknown product in history')
        key = (product, str(release_id))
        if kind == 'release':
            if key in state['releases']:
                raise ValueError('Duplicate release identity in history')
            state['releases'][key] = {'tag': tag, 'published': published or ''}
        elif kind == 'asset':
            values = list(row[9:9 + len(state['runs'])])
            if not all(value is None or type(value) is int and value >= 0 for value in values):
                raise ValueError('Invalid historical download count')
            asset_key = (*key, str(asset_id))
            if asset_key in state['assets'] or group not in GROUPS or platform not in PLATFORMS:
                raise ValueError('Invalid or duplicate artifact identity in history')
            state['assets'][asset_key] = dict(name=name, group=group, platform=platform, values=values)
        else:
            raise ValueError('Unknown record in history')
    if any(key[:2] not in state['releases'] for key in state['assets']):
        raise ValueError('Artifact refers to a missing release')
    return state


def load_history(path):
    if path is None:
        return empty_history()
    try:
        workbook = load_workbook(path, data_only=False)
        try:
            if METADATA not in workbook.sheetnames:
                raise HistoryError('The existing workbook uses an old format; start with a new document')
            return read_metadata(workbook[METADATA])
        finally:
            workbook.close()
    except HistoryError:
        raise
    except Exception as exc:
        raise HistoryError(f'Existing history cannot be read ({type(exc).__name__}); report will not be overwritten') from exc


def write_total(sheet, row, label, references, values, styles):
    sheet.write_string(row, 0, label, styles['total'])
    for index, (cells, numbers) in enumerate(zip(references, values), 1):
        recorded = [number for number in numbers if number is not None]
        cached = sum(recorded) if recorded else '—'
        formula = f'=IF(COUNT({cells})=0,"—",SUM({cells}))'
        sheet.write_formula(row, index, formula, styles['total'], cached)


def write_product(book, product, state, styles):
    sheet = book.add_worksheet(product)
    sheet.hide_gridlines(2)
    sheet.freeze_panes(4, 1)
    sheet.set_default_row(23)
    sheet.set_row(3, 36)
    sheet.set_column(0, 0, 64)
    sheet.set_column(1, len(state['runs']), 22)
    sheet.merge_range(0, 0, 1, max(1, len(state['runs'])), product, styles['title'])
    sheet.write_string(2, 0, 'Lifetime downloads · tracked stable releases', styles['note'])
    sheet.write_row(3, 0, ['Release / artifact', *state['runs']], styles['header'])
    releases = sorted((key for key in state['releases'] if key[0] == product),
                      key=lambda key: (state['releases'][key]['published'], state['releases'][key]['tag']), reverse=True)
    row = 4
    for release_key in releases:
        release = state['releases'][release_key]
        sheet.merge_range(row, 0, row, max(1, len(state['runs'])), release['tag'], styles['release'])
        row += 1
        entries = [(key, entry) for key, entry in state['assets'].items() if key[:2] == release_key]
        duplicates = Counter(entry['name'] for _, entry in entries)
        for group in GROUPS:
            group_rows, group_values = [], [[] for _ in state['runs']]
            if not any(entry['group'] == group for _, entry in entries):
                continue
            sheet.merge_range(row, 0, row, max(1, len(state['runs'])), group, styles['group'])
            row += 1
            for platform in PLATFORMS:
                assets = sorted(((key, entry) for key, entry in entries
                                 if entry['group'] == group and entry['platform'] == platform),
                                key=lambda pair: (pair[1]['name'].lower(), pair[0][-1]))
                if not assets:
                    continue
                sheet.merge_range(row, 0, row, max(1, len(state['runs'])), platform, styles['platform'])
                row += 1
                start = row
                for key, entry in assets:
                    label = entry['name'] if duplicates[entry['name']] == 1 else f"{entry['name']} [asset {key[-1]}]"
                    sheet.write_string(row, 0, label, styles['text'])
                    for column, value in enumerate(entry['values'], 1):
                        sheet.write(row, column, value if value is not None else '—', styles['number'])
                    row += 1
                references = [f'{xl_rowcol_to_cell(start, col)}:{xl_rowcol_to_cell(row - 1, col)}'
                              for col in range(1, len(state['runs']) + 1)]
                values = [[entry['values'][index] for _, entry in assets] for index in range(len(state['runs']))]
                write_total(sheet, row, f'Total {platform}', references, values, styles)
                group_rows.append(row)
                for index, numbers in enumerate(values):
                    group_values[index].extend(numbers)
                row += 1
            references = [','.join(xl_rowcol_to_cell(total_row, col) for total_row in group_rows)
                          for col in range(1, len(state['runs']) + 1)]
            write_total(sheet, row, f'Total — {group.lower()}', references, group_values, styles)
            row += 1
        if not entries:
            sheet.write_string(row, 0, 'No attached artifacts', styles['note'])
            row += 1
        row += 1
    sheet.set_landscape()
    sheet.print_area(0, 0, row - 1, len(state['runs']))
    sheet.repeat_rows(3)
    return sheet



def write_docker(book, state, styles):
    sheet = book.add_worksheet('Docker Docs')
    sheet.hide_gridlines(2)
    sheet.freeze_panes(4, 1)
    sheet.set_default_row(25)
    sheet.set_row(3, 36)
    sheet.set_column(0, 0, 64)
    sheet.set_column(1, len(state['runs']), 22)
    last = max(1, len(state['runs']))
    sheet.merge_range(0, 0, 1, last, 'Docker Docs', styles['title'])
    sheet.merge_range(2, 0, 2, last, f'Docker Hub / {DOCKER_REPOSITORY}', styles['note'])
    sheet.write_row(3, 0, ['Metric', *state['runs']], styles['header'])
    sheet.write_string(4, 0, 'Total pulls (lifetime)', styles['text'])
    sheet.write_string(5, 0, 'Increase since previous observation', styles['total'])
    for column, value in enumerate(state['docker'], 1):
        sheet.write(4, column, value if value is not None else '—', styles['number'])
        previous = state['docker'][column - 2] if column > 1 else None
        if value is None or previous is None or value < previous:
            sheet.write_string(5, column, '—', styles['total'])
        else:
            current_cell = xl_rowcol_to_cell(4, column)
            previous_cell = xl_rowcol_to_cell(4, column - 1)
            sheet.write_formula(5, column, f'={current_cell}-{previous_cell}',
                                styles['total'], value - previous)
    sheet.merge_range(7, 0, 7, last, 'Pulls are not installations or unique users.', styles['note'])
    sheet.merge_range(8, 0, 8, last, 'First observation or decreased counter: no increase value.', styles['note'])
    sheet.set_landscape()
    sheet.print_area(0, 0, 8, len(state['runs']))


def save_history(path, state):
    path.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary = tempfile.mkstemp(suffix='.xlsx', dir=path.parent)
    os.close(handle)
    try:
        with xlsxwriter.Workbook(temporary, {'strings_to_formulas': False, 'strings_to_urls': False}) as book:
            base = {'font_name': 'Calibri', 'font_size': 11, 'font_color': '#243746', 'valign': 'vcenter'}
            definitions = {
                'title': dict(bold=True, font_size=22, font_color='#17365D'),
                'note': dict(font_size=10, font_color='#64748B'),
                'header': dict(bold=True, bg_color='#24476B', font_color='#FFFFFF', text_wrap=True),
                'release': dict(bold=True, font_size=13, bg_color='#24476B', font_color='#FFFFFF'),
                'group': dict(bold=True, bg_color='#DCE6EF'),
                'platform': dict(bold=True, bg_color='#F0F4F8'),
                'text': {}, 'number': dict(num_format='#,##0', align='right'),
                'total': dict(bold=True, top=1, top_color='#CBD5E1', num_format='#,##0', bg_color='#E8EFF5')}
            styles = {key: book.add_format(dict(base, **value)) for key, value in definitions.items()}
            for product in PRODUCTS:
                write_product(book, product, state, styles)
            write_docker(book, state, styles)
            sheet = book.add_worksheet(METADATA)
            sheet.hide()
            sheet.write_string(0, 0, SCHEMA)
            sheet.write_row(1, 0, ['Kind', 'Product', 'Release ID', 'Asset ID', 'Tag', 'Published',
                                   'Filename', 'Group', 'Platform', *state['runs']])
            row = 2
            for key, release in state['releases'].items():
                sheet.write_row(row, 0, ['release', *key, None, release['tag'], release['published']])
                row += 1
            for key, entry in state['assets'].items():
                sheet.write_row(row, 0, ['asset', *key, None, None, entry['name'], entry['group'],
                                       entry['platform'], *entry['values']])
                row += 1
            sheet.write_row(row, 0, ['docker', 'Docker Docs', None, None, None, None,
                                    DOCKER_REPOSITORY, None, None, *state['docker']])
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
